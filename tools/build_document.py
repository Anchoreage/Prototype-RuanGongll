from pathlib import Path
import sys, math, re
sys.stdout.reconfigure(encoding='utf-8')
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from document_fonts import CJK_FONT, LATIN_FONT, rewrite_docx

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'deliverables'
QA=ROOT/'.qa'
OUT.mkdir(exist_ok=True);QA.mkdir(exist_ok=True)
FONT=r'C:\Windows\Fonts\msyh.ttc'
TIMES_FONT=r'C:\Windows\Fonts\times.ttf'
def font(n):return ImageFont.truetype(FONT,n)

# Deterministic UML use case drawing, independent of the interface design image.
W,H=2400,1330
im=Image.new('RGB',(W,H),'white');d=ImageDraw.Draw(im)
ink='#233a31';muted='#62796c';green='#ecf4ef'
def text_parts(line,size):
 return [(part,ImageFont.truetype(TIMES_FONT,size) if ord(part[0])<256 else font(size)) for part in re.findall(r'[\x00-\xff]+|[^\x00-\xff]+',line)]

def text_width(line,size):
 return sum(d.textlength(part,font=f) for part,f in text_parts(line,size))

def txt(x,y,text,size=31,fill=ink):
 for n,line in enumerate(text.split('\n')):
  left=x-text_width(line,size)/2
  for part,f in text_parts(line,size):
   d.text((left,y+n*(size+8)),part,font=f,fill=fill)
   left+=d.textlength(part,font=f)
def actor(x,y,label):
 d.ellipse((x-19,y,x+19,y+38),outline=ink,width=3)
 d.line([(x,y+38),(x,y+108)],fill=ink,width=3)
 d.line([(x-45,y+61),(x+45,y+61)],fill=ink,width=3)
 d.line([(x,y+108),(x-35,y+155)],fill=ink,width=3)
 d.line([(x,y+108),(x+35,y+155)],fill=ink,width=3)
 txt(x,y+176,label,31)
def oval(x,y,label,learning=False,width=350,height=102):
 d.ellipse((x-width/2,y-height/2,x+width/2,y+height/2),outline=ink,width=3,fill=green if learning else 'white')
 lines=label.split('\n');txt(x,y-(len(lines)*(31+8))/2-3,label,31)
def association(points):d.line(points,fill=ink,width=3)
def arrow(points,label,at):
 for a,b in zip(points,points[1:]):
  dx,dy=b[0]-a[0],b[1]-a[1];length=math.hypot(dx,dy)
  for t in range(0,int(length),16):
   t2=min(t+9,length);d.line((a[0]+dx*t/length,a[1]+dy*t/length,a[0]+dx*t2/length,a[1]+dy*t2/length),fill=muted,width=2)
 a,b=points[-2:];ang=math.atan2(b[1]-a[1],b[0]-a[0]);sz=14
 d.line([(b[0]-sz*math.cos(ang-.45),b[1]-sz*math.sin(ang-.45)),b,(b[0]-sz*math.cos(ang+.45),b[1]-sz*math.sin(ang+.45))],fill=muted,width=2)
 width=text_width(label,25);d.rectangle((at[0]-width/2-7,at[1]-1,at[0]+width/2+7,at[1]+32),fill='white');txt(*at,label,25,muted)

d.rectangle((260,45,2135,1235),outline=ink,width=3)
txt(1190,65,'法律论证标注与教学系统',40)
actor(115,380,'学生\n标注者');actor(2270,200,'教师\n任务创建者');actor(2270,850,'裁定者\n教师或专家')
# Student goals on left, shared mandatory behavior in middle, privileged goals right.
left=[(455,230,'学习标注指南',True),(455,415,'命题标注',False),(455,600,'关系标注',False),(455,785,'独立完成练习\n并提交',True),(455,990,'查看学习反馈\n扩展点 发现差异',True)]
for x,y,l,g in left:oval(x,y,l,g,340,112)
middle=[(1050,230,'校验并预览指南\n记录版本',False),(1050,505,'更新论证图示',True),(1050,990,'对照标准或\n裁定结果',True)]
for x,y,l,g in middle:oval(x,y,l,g,350,112)
right=[(1700,230,'从文件导入\n标注指南',False),(1700,415,'配置练习与\n标准样例',True),(1700,670,'裁定关系的标注',False),(1650,860,'查看独立版本',False),(1650,1030,'记录理由并\n保存最终版本',False)]
for x,y,l,g in right:oval(x,y,l,g,365,112)
oval(720,1170,'反思并重新练习',True,360,92)
# Plain associations deliberately routed along the system boundary.
association([(160,441),(275,441),(275,230),(285,230)])
for y in [415,600,785,990]:association([(160,441),(275,441),(275,y),(285,y)])
for y in [230,415]:association([(2225,261),(2110,261),(2110,y),(1883,y)])
association([(2225,911),(2110,911),(2110,670),(1883,670)])
# include arrows go from the base use case to the mandatory included behavior.
arrow([(1517,230),(1225,230)],'«include»',(1370,192))
arrow([(625,415),(790,415),(790,480),(875,480)],'«include»',(743,376))
arrow([(625,600),(790,600),(790,530),(875,530)],'«include»',(743,610))
arrow([(625,990),(875,990)],'«include»',(751,953))
arrow([(1645,722),(1510,790),(1510,825)],'«include»',(1480,746))
arrow([(1755,724),(1905,760),(1905,1030),(1833,1030)],'«include»',(1930,850))
# extend points to the base use case; no implied process arrows between use cases.
arrow([(540,1170),(455,1170),(455,1046)],'«extend»',(390,1090))
txt(1220,1175,'扩展条件  已查看反馈且存在需修正的差异',25,muted)
txt(1200,1270,'实线  参与者关联    虚线箭头  包含或扩展    浅绿用例  支持学习目标',28,muted)
im.save(QA/'system-use-cases.png')

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 s=doc.styles[name];s.font.name=LATIN_FONT;s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),CJK_FONT);s.font.color.rgb=RGBColor(0,0,0)
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[name];st.font.italic=False;st.font.underline=False
 for border in st._element.xpath('.//w:pBdr'):border.getparent().remove(border)
normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.3;normal.paragraph_format.space_after=Pt(7)
for name,n in [('Title',26),('Subtitle',12),('Heading 1',17),('Heading 2',12),('Caption',9)]:doc.styles[name].font.size=Pt(n)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(3);doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(8);doc.styles['Heading 2'].paragraph_format.space_after=Pt(7)
doc.core_properties.title='法律论证标注系统用例设计与项目原型';doc.core_properties.subject='系统用例图和三个核心用例描述';doc.core_properties.author='';doc.core_properties.keywords='法律论证 标注 用例 教学 原型'
def p(text,style=None):return doc.add_paragraph(text,style)
def heading(text):doc.add_heading(text,level=1)
def table(headers,rows,widths,fontsize=9.3):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 pr.append(borders)
 for vals in [headers]+rows:
  row=t.rows[0] if vals is headers else t.add_row()
  no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
  ishead=vals is headers
  if ishead:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for j,(cell,val,w) in enumerate(zip(row.cells,vals,widths)):
   cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cell.text=str(val);tcpr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for edge,v in [('top',60),('bottom',60),('left',110),('right',110)]:el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(v));el.set(qn('w:type'),'dxa');m.append(el)
   tcpr.append(m);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'284D41' if ishead else ('F1F5F3' if len(t.rows)%2==0 else 'FFFFFF'));tcpr.append(shade)
   for para in cell.paragraphs:
    para.paragraph_format.line_spacing=1.16;para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1)
    if j==0 and len(headers)>2:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:run.font.name=LATIN_FONT;run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),CJK_FONT);run.font.size=Pt(fontsize);run.bold=ishead;run.font.color.rgb=RGBColor.from_string('FFFFFF' if ishead else '27362D')
 spacer=p('');spacer.paragraph_format.space_after=Pt(4);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=1;spacer.add_run().font.size=Pt(2);spacer.paragraph_format.keep_with_next=False
 return t
def newpage(title):
 h=doc.add_heading(title,level=1);h.paragraph_format.page_break_before=True
def footer(s):
 f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 r=f.add_run('法律论证标注系统    ');r.font.name=LATIN_FONT;r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),CJK_FONT);r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('748076')
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
footer(sec)

doc.add_paragraph('法律论证标注系统',style='Title')
p('系统用例设计与项目原型',style='Subtitle')
p('课程项目交付文档    2026年10月7日')
heading('1 业务目标与系统边界')
p('系统围绕裁判文书的说理部分，支持指南配置、命题与关系标注、关系裁定及结果反馈。教学目标是帮助学生更快识别事实与规范，更好理解从前提到结论的论证结构。学生先独立判断，提交后再对照标准或裁定结果，并通过差异解释进行反思练习。')
table(['业务目标','系统设计','可观察的验收指标'],[
 ['更快掌握框架','就地查看标签定义与示例；原文、命题列表和图示联动；按步骤完成练习。','记录同等难度独立练习用时；比较前后测完成时间。'],
 ['更好掌握框架','练习提交后比较边界、标签和关系；展示教师裁定理由；允许新建练习纠错。','比较命题边界与标签一致率、关系结构完整率，以及相同错误再次出现的比例。'],
 ['保证学习可复核','任务绑定指南版本；保留学生原始提交、标准版本和裁定记录。','任一反馈可追溯至原文范围、规则版本及最终关系。'],
],[3.0,7.1,6.9],9.5)
doc.add_heading('参与者与职责',level=2)
table(['参与者','主要职责'],[
 ['学生或标注者','阅读指南，独立进行命题与关系标注，提交练习，查看差异与解释。'],
 ['教师或任务创建者','导入并确认指南，配置任务和标准样例，维护教学材料。'],
 ['裁定者','教师或专家人工比对独立版本，采用或重标关系，记录理由并定稿。'],
],[4.1,12.9],9.5)
p('范围说明：命题性质使用 SF、GF、SM、GM；关系使用 S、A、J、M、I。“裁定关系的标注”是对标注结果进行复核定稿的用例。合同及合同解释按指南正文使用 SM-C。任务创建、标注、裁定和输出分别控制编辑权限；学生练习使用独立入口与临时数据。')
p('学习指标属于后续教学评价方案，原型展示计算与反馈方式，实际学习成效需通过同等难度的前后测验证。')

s=doc.add_section(WD_SECTION_START.NEW_PAGE);s.orientation=WD_ORIENT.LANDSCAPE;s.page_width=Cm(29.7);s.page_height=Cm(21);s.top_margin=Cm(1.5);s.bottom_margin=Cm(1.4);s.left_margin=Cm(1.7);s.right_margin=Cm(1.7)
heading('2 系统用例图')
para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.line_spacing=1;para.paragraph_format.space_after=Pt(0);run=para.add_run();run.add_picture(str(QA/'system-use-cases.png'),width=Cm(24.6));run._r.xpath('.//wp:docPr')[0].set('descr','系统用例图：学生、教师及裁定者位于系统边界外，关联标注、指南、练习反馈及裁定用例；包含与扩展关系以虚线箭头表示。')
p('图1  系统用例图',style='Caption')
explanation=p('«include» 表示每次执行基础用例都必须执行的子用例；«extend» 从可选扩展指向基础用例。“反思并重新练习”在查看反馈且发现差异时扩展“查看学习反馈”。图中的连接表达参与者和功能关系，不表示业务执行先后顺序。');explanation.paragraph_format.line_spacing=1;explanation.paragraph_format.space_after=Pt(0)
s=doc.add_section(WD_SECTION_START.NEW_PAGE);s.orientation=WD_ORIENT.PORTRAIT;s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.9);s.bottom_margin=Cm(1.7);s.left_margin=Cm(2);s.right_margin=Cm(2)

heading('3 从文件导入标注指南')
table(['用例字段','描述'],[
 ['编号与级别','UC-01  从文件导入标注指南；用户目标级；优先级高。'],
 ['目标与参与者','教师或任务创建者将指南文件导入系统，确认规则并形成可追溯的指南版本。'],
 ['触发事件','在“标注指南”页选择文件或拖入文件。'],
 ['前置条件','用户具有指南配置权限；处于任务准备阶段或正在创建新任务；文件可读取。'],
 ['成功后置条件','保存指南名称、版本、来源、标签定义与关系规则；新任务绑定该版本；学生可阅读。'],
 ['最小保证','校验失败、用户取消或保存失败时，不替换已启用指南，不修改原任务标注。'],
 ['输入与输出','输入指南文件；输出导入预览、问题提示和可启用的新版本。'],
],[3.0,14.0])
doc.add_heading('主成功场景',level=2)
table(['步骤','参与者动作','系统响应'],[
 ['1','打开指南页并选择导入。','提示支持格式、大小限制、字段要求和当前指南版本。'],
 ['2','选择文件并确认读取。','检查格式及大小，解析JSON结构或提取PDF、DOCX、TXT正文。'],
 ['3','核对名称、版本与正文。','展示来源、候选规则和定义；非结构化文件明确提示人工确认。'],
 ['4','修正候选定义与版本信息。','检验必填字段、唯一代码和本框架所需的命题与关系类型。'],
 ['5','确认规则并启用。','说明任务绑定及新建练习影响，避免在进行中的任务内替换标签。'],
 ['6','确认新任务使用该版本。','保存新版本与来源，创建绑定该指南的练习，展示成功状态。'],
],[1.2,6.1,9.7])

newpage('3 从文件导入标注指南  续')
doc.add_heading('替代与异常流程',level=2)
table(['编号','对应步骤与条件','系统处理','结果'],[
 ['A1','步骤2  PDF、DOCX或TXT','仅提取正文；用课程框架生成规则候选；教师逐项核对后确认，不能把正文直接视为已解析规则。','转步骤3。'],
 ['A2','步骤4  需要修改候选','允许调整名称、版本和定义；复杂结构先编辑JSON模板后重新导入。','重新校验。'],
 ['E1','步骤2  不支持格式或超过8 MB','说明具体限制，保留当前指南与任务。','重新选择或取消。'],
 ['E2','步骤2或4  文件损坏、代码重复或字段缺失','显示明确问题，禁止启用，原版本继续生效。','修正后重试。'],
 ['E3','步骤2  扫描PDF无可提取文字','提示先进行文字识别或使用结构化模板。','本次导入失败。'],
 ['E4','步骤5或6  取消或保存失败','取消候选；保存失败不得显示已启用。','不产生新有效版本。'],
],[1.2,4.2,8.8,2.8])
doc.add_heading('规则与数据要求',level=2)
table(['项目','规定'],[
 ['导入边界','支持 .json、文字型 .pdf、.docx、UTF-8 .txt；原型文件大小上限8 MB，DOCX正文上限4 MB。'],
 ['标签校验','SF、GF、SM、GM均有名称和定义；S、A、J、M、I均有关系定义；代码唯一；子类型须属于已支持集合。'],
 ['版本规则','以指南正文的版本声明为准；任务开始后规则锁定，指南更新创建新任务，不覆盖旧任务记录。'],
 ['记录字段','指南ID、名称、版本、文件名、导入时间、标签及关系定义、原文或来源引用、确认人。'],
 ['教学价值','学生可就地查阅统一规则与示例，避免在不同版本之间反复查找和混用概念。'],
],[3.0,14.0])
doc.add_heading('验收示例',level=2)
p('导入模板JSON，出现完整规则预览；教师确认后新练习显示同一版本。将两项标签改为相同代码，系统明确提示代码重复并禁止启用。导入课程PDF，正文可查阅且必须经过规则候选确认。')

newpage('4 命题标注')
table(['用例字段','描述'],[
 ['编号与级别','UC-02  命题标注；用户目标级；优先级高。'],
 ['目标与参与者','学生或标注者识别文本中完整判断的边界与性质，形成与原文一致的命题记录。'],
 ['触发事件','在原文中拖选文字，或选择命题列表中的已有命题。'],
 ['前置条件','已获得原文与有效指南；具有本任务标注权限；任务尚未提交或定稿。'],
 ['成功后置条件','保存原文范围、标签与稳定标识；命题按原文顺序编号；高亮、列表和图示同步更新。'],
 ['最小保证','无效选区或保存失败不产生命题；已提交版本不被覆盖；撤销可回到上一有效状态。'],
 ['输入与输出','输入选区及一级、子类型标签；输出命题记录、原文高亮及对应矩形节点。'],
],[3.0,14.0])
doc.add_heading('主成功场景',level=2)
table(['步骤','参与者动作','系统响应'],[
 ['1','阅读裁判理由并选择一个完整判断。','计算原文起止索引，展示选中文字和字符数。'],
 ['2','查看类型定义及示例。','就地显示SF、GF、SM、GM区分规则和判断提示。'],
 ['3','选择一级标签。','显示该标签定义；需要细分时提供对应子类型。'],
 ['4','选择GM子类型或SM-C等适用子类型。','根据指南限制可选值，避免重复标签与不适用子类型。'],
 ['5','保存命题。','校验选区不为空、范围有效且无重复或重叠，检查既有关系的约束。'],
 ['6','查看保存结果。','生成稳定ID，按原文位置显示p1、p2等编号，同步高亮、完整文本列表和节点。'],
 ['7','继续标注或提交练习。','保存草稿；提交后锁定本次结果，转入学习反馈。'],
],[1.2,6.1,9.7])

newpage('4 命题标注  续')
doc.add_heading('替代与异常流程',level=2)
table(['编号','对应步骤与条件','系统处理','结果'],[
 ['A1','步骤1  修改已有命题','点击列表或原文高亮查看全文；修改类型并校验依赖关系。','保存后同步更新。'],
 ['A2','步骤6  删除或调整边界','删除命题并提示关联关系处理；边界调整通过删除后重新选择实现。','依赖关系一并清理。'],
 ['A3','步骤6  误操作','草稿阶段支持撤销、重做；依赖关系与稳定ID随状态一起恢复。','恢复上一有效版本。'],
 ['E1','步骤1或5  空选区、越界、重复或重叠','解释问题并禁止保存；保留当前有效标注。','调整选区后重试。'],
 ['E2','步骤5  类型修改导致匹配失效','指出受影响关系；先修改或删除关系，再调整命题标签。','不保存无效状态。'],
 ['E3','任一步骤  已提交或无编辑权限','标注控制只读；可回看原始结果或开启新练习。','不覆盖提交版本。'],
],[1.2,4.2,8.8,2.8])
doc.add_heading('规则与数据要求',level=2)
table(['项目','规定'],[
 ['边界规则','以显性文本中的完整判断为单位；去除首尾空白；同一练习不接受重叠命题。'],
 ['类型规则','事实与规范、一般与个别分别判断。劳务合同关系的法律定性使用SM；履行事实使用SF。'],
 ['编号与标识','展示序号按起始位置排序并动态重排；关系引用稳定ID，不能引用易变化的展示序号。'],
 ['原文定位','记录[start,end)字符范围；原型使用UTF-16索引。原文变更需重新校正范围或创建新任务。'],
 ['记录字段','任务ID、原文标识、命题ID、起止索引、原文、类型及子类型、指南版本、标注者与时间。'],
 ['教学价值','在原文与图示间双向定位，帮助学生把文字判断转化为可见论证要素；提交后反馈边界与标签差异。'],
],[3.0,14.0])
doc.add_heading('验收示例',level=2)
p('选择“原告已依约提供劳务，”并标为SF，列表显示完整文字且原文高亮。先后选择同一范围时不生成重复命题。删除依赖于关系的命题后不留下悬空引用；撤销恢复命题与相应关系。')

newpage('5 裁定关系的标注')
table(['用例字段','描述'],[
 ['编号与级别','UC-03  裁定关系的标注；用户目标级；优先级高。'],
 ['目标与参与者','裁定者对同一文书的独立关系版本进行人工比对，形成有依据的最终关系版本。'],
 ['触发事件','裁定者打开已提交的关系标注，选择一条待复核关系。'],
 ['前置条件','具有裁定权限；原文与指南版本已确定；至少存在一个已提交版本；相关命题可定位。'],
 ['成功后置条件','保留原始版本；保存最终关系、理由、裁定者与时间；结果锁定并可导出，学生可查阅解释。'],
 ['最小保证','取消或校验失败时不定稿；任一原始标注版本不得被裁定操作覆盖。'],
 ['输入与输出','输入原文、命题与独立关系版本；输出最终结构、裁定理由及可追溯记录。'],
],[3.0,14.0])
doc.add_heading('主成功场景',level=2)
table(['步骤','参与者动作','系统响应'],[
 ['1','打开关系裁定工作区并选择条目。','展示原文、命题编号、指南版本和独立标注版本。'],
 ['2','对照原文与规则人工比较版本。','并排显示完整表达式，便于检查类型、成员、方向和嵌套。'],
 ['3','采用某一版本，或重新标注最终关系。','生成最终候选；重标时可输入合法的嵌套关系表达式。'],
 ['4','填写裁定理由。','要求解释规范依据及差异原因，作为学生可理解的教学反馈。'],
 ['5','确认裁定并定稿。','校验命题存在、参数数量、匹配类型、嵌套结构及理由非空。'],
 ['6','查看最终结果。','保存候选、来源版本、裁定理由及时间，保持原始版本不变并锁定最终结果。'],
 ['7','导出或查看反馈。','全部条目定稿后导出最终结果；对照学生结果时显示差异与理由。'],
],[1.2,6.1,9.7])

newpage('5 裁定关系的标注  续')
doc.add_heading('替代与异常流程',level=2)
table(['编号','对应步骤与条件','系统处理','结果'],[
 ['A1','步骤1  只有一个版本','仍展示该版本供复核；不因没有分歧而跳过最终确认。','进入步骤2。'],
 ['A2','步骤3  两个版本均不适合','选择重新标注，创建有效结构并说明原因。','重标版本作为候选。'],
 ['A3','步骤2  难以判断','正式系统保存待讨论状态，专家线下讨论后由裁定者记录理由并确认。','暂不定稿。'],
 ['E1','步骤5  起终点缺失、方向或嵌套不合法','指出具体错误；保持候选与原始版本。','修正后重新确认。'],
 ['E2','步骤5  理由为空或保存失败','禁止定稿，不更新已生效结果，提示补充或重试。','条目保持待裁定。'],
 ['E3','步骤6后  请求修改最终结果','只读显示；纠错需新任务或正式系统的受控修订流程。','保留原定稿记录。'],
],[1.2,4.2,8.8,2.8])
doc.add_heading('规则与数据要求',level=2)
table(['项目','规定'],[
 ['关系规则','S、A、M具有方向；J、I至少两个成员。M由个别判断指向一般判断。A可指向命题或已有支持关系。'],
 ['嵌套规则','支持S(J(...),p)、M(J(...),p)、S(M(...),p)、A(p,S(...))等；逐层保存关系单元，不省略中间结构。'],
 ['裁定原则','依据文本所声称的逻辑联系和已绑定指南进行复核；裁定理由解释标注选择，不能只给出“正确”结论。'],
 ['记录字段','任务与条目ID、原始版本引用、最终关系、采用或重标方式、理由、指南版本、裁定者、时间与状态。'],
 ['教学价值','学生从不同关系结构及裁定理由理解前提、结论、方向与组合之间的差别，再进行独立重练。'],
],[3.0,14.0])
doc.add_heading('验收示例',level=2)
p('当A版本为S(p6,p5)、B版本为S(p5,p6)时，裁定者采用从证据判断指向欠款事实的版本，填写理由后定稿。裁定结果记录最终结构和来源，双方原始版本仍可查看；理由为空时不得定稿。')

newpage('6 项目原型与演示说明')
p('原型采用“原文阅读—命题识别—关系构建—提交反馈”的学习工作台。教师通过同一原型的演示角色入口导入指南、核对规则，并在裁定页比较固定教学案例。')
shot=OUT/'原型预览.png'
if shot.exists():
 para=doc.add_paragraph();r=para.add_run();r.add_picture(str(shot),width=Cm(17));r._r.xpath('.//wp:docPr')[0].set('descr','法律论证标注教学原型：原文、命题列表、学习助手和论证图示。');p('图2  原型学习工作台',style='Caption')
table(['演示步骤','操作与预期结果'],[
 ['导入指南','切换教师 / 裁定者，打开标注指南；导入JSON模板或课程PDF；核对候选后创建新练习。'],
 ['命题标注','在裁判理由中拖选文字，选择类型并保存；列表、高亮和节点同时更新；可撤销与重做。'],
 ['关系构建','选择S、A、J、M、I及有效对象；先创建组合，再构建匹配或支持；查看完整嵌套表达式。'],
 ['裁定关系','打开关系裁定；采用A或B版本，或重标；填写理由后定稿；完成后导出最终版本。'],
 ['查看反馈','学生提交练习后分别对照命题边界、标签和关系结构；查看解释、用时并重新独立练习。'],
],[3.0,14.0],9.2)
newpage('7 实现范围与资料依据')
p('已实现：指南文件读取与预览、命题标注和依赖清理、五种关系及嵌套、实时图示、撤销重做、本机草稿、JSON导出、固定案例裁定、练习对照反馈，以及桌面和手机布局。')
p('原型采用浏览器本地存储及演示角色切换；正式多人协作、服务器权限校验、任务批量分配、真实多人版本对齐、受控修订和Excel导出属于后续开发范围。导出的JSON包含指南版本与字符索引。')
doc.add_heading('资料依据',level=2)
p('1  裁判文书标记需求_按阶段组织（2）.docx，文档版本v1.1，整理日期2026年4月14日。\n2  中文裁判文书说理部分法律论证结构标注与图示指南，重点参照命题分类、五种关系、图示规则、标注示例及培训反馈要求。压缩包文件名标为v1.2，正文声明为v1.1；本原型按正文标识版本。')
doc.save(OUT/'法律论证标注系统用例设计与项目原型.docx')
rewrite_docx(OUT/'法律论证标注系统用例设计与项目原型.docx')
print(OUT/'法律论证标注系统用例设计与项目原型.docx')
